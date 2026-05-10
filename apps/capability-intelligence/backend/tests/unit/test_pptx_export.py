"""PPTX renderer — emits a valid pptx blob with expected slide count."""

import io

from pptx import Presentation

from app.services import pptx_export


def _digest_fixture(n_priorities: int = 3) -> dict:
    priorities = [
        {
            "sub_cap_id": f"P1C1.1.{i + 1}",
            "sub_cap_name": f"Priority {i + 1}",
            "state": "RISING",
            "score": 78.5 - i * 5,
            "confidence": 0.75,
            "narrative": "Lorem ipsum sample narrative covering this subcap.",
            "recommendation": "Recommend a follow-up review next quarter.",
            "evidence_sows": [{
                "sow_id": "sow-x", "client": "Wells Fargo", "status": "active",
                "excerpt": "active engagement covering this subcap",
                "method": "exact_id", "confidence": 99,
            }],
            "evidence_benchmarks": [{
                "metric_id": "tech_spend_pct_revenue", "cohort_id": "us_banks_gsib",
                "period": "2025-Q4", "verdict": "INDICATIVE", "n": 3,
                "p25": 11.0, "p50": 11.6, "p75": 12.0,
            }],
            "evidence_news": [{
                "id": "news-1", "title": "FS GenAI strategy",
                "source": "americanbanker.com", "published_at": "2026-04-12",
                "kind": "news", "url": None,
            }],
            "delta": {"previous_state": "EMERGING", "previous_period": "2026-Q1"},
            "chain_id": "chain-abc12345",
            "cost_usd": 0.0,
        }
        for i in range(n_priorities)
    ]
    return {
        "digest_id": "digest-retail-banking-2026-Q2",
        "subvertical": "retail-banking",
        "period": "2026-Q2",
        "previous_period": "2026-Q1",
        "generated_at": "2026-05-10T12:00:00+00:00",
        "model": "opus",
        "priorities": priorities,
        "summary": "Q2 2026 digest summary paragraph.",
        "sources_count": 9,
        "total_cost_usd": 0.0,
    }


def test_render_returns_valid_pptx_bytes():
    n = 3
    blob = pptx_export.render(_digest_fixture(n))
    assert blob[:2] == b"PK"  # ZIP / OOXML magic
    prs = Presentation(io.BytesIO(blob))
    # 1 title + 1 summary + N priorities + 1 watchlist
    assert len(prs.slides) == 1 + 1 + n + 1


def test_render_handles_zero_priorities():
    digest = _digest_fixture(0)
    blob = pptx_export.render(digest)
    prs = Presentation(io.BytesIO(blob))
    # 1 title + 1 summary + 0 priorities + 1 watchlist
    assert len(prs.slides) == 3


def test_render_includes_subvertical_in_title_slide():
    blob = pptx_export.render(_digest_fixture(2))
    prs = Presentation(io.BytesIO(blob))
    title_text = []
    for shape in prs.slides[0].shapes:
        if shape.has_text_frame:
            title_text.append(shape.text_frame.text)
    joined = "\n".join(title_text)
    assert "retail-banking" in joined
    assert "2026-Q2" in joined


def test_render_priority_slide_text_includes_state_and_name():
    blob = pptx_export.render(_digest_fixture(1))
    prs = Presentation(io.BytesIO(blob))
    # priority slide is index 2 (after title + summary)
    priority_slide = prs.slides[2]
    text_blob = "\n".join(
        s.text_frame.text for s in priority_slide.shapes if s.has_text_frame
    )
    assert "RISING" in text_blob
    assert "Priority 1" in text_blob
    # Q-over-Q delta string renders
    assert "EMERGING" in text_blob
