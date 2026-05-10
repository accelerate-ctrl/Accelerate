"""Text extraction from SOW files.

Local mode (no GCP creds): pypdf for .pdf, python-docx for .docx, raw read
for .txt. GCP mode (Document AI) is wired through the same `extract()` API
in a later batch — switch on `Settings.use_gcp` here when the DocAI
processor is configured.
"""
from __future__ import annotations

import logging
from dataclasses import dataclass
from pathlib import Path

log = logging.getLogger(__name__)

SUPPORTED_EXTS = {".pdf", ".docx", ".txt", ".md"}


@dataclass
class ExtractedDocument:
    text: str
    page_count: int
    char_count: int
    extractor: str  # "txt" | "pypdf" | "docx" | "documentai"


def extract(content: bytes, *, file_name: str) -> ExtractedDocument:
    ext = Path(file_name).suffix.lower()
    if ext == ".txt" or ext == ".md":
        text = content.decode("utf-8", errors="replace")
        return ExtractedDocument(text=text, page_count=1, char_count=len(text), extractor="txt")
    if ext == ".pdf":
        return _extract_pdf(content)
    if ext == ".docx":
        return _extract_docx(content)
    raise ValueError(f"unsupported extension: {ext}")


def _extract_pdf(content: bytes) -> ExtractedDocument:
    from io import BytesIO
    from pypdf import PdfReader
    reader = PdfReader(BytesIO(content))
    parts: list[str] = []
    for page in reader.pages:
        parts.append(page.extract_text() or "")
    text = "\n\n".join(parts)
    return ExtractedDocument(text=text, page_count=len(reader.pages), char_count=len(text), extractor="pypdf")


def _extract_docx(content: bytes) -> ExtractedDocument:
    from io import BytesIO
    from docx import Document
    doc = Document(BytesIO(content))
    text = "\n".join(p.text for p in doc.paragraphs if p.text)
    return ExtractedDocument(text=text, page_count=1, char_count=len(text), extractor="docx")
