"""Multi-format document ingestion: BRDs and SDDs arrive as .docx, .pdf,
.html or plain text/markdown — the app recognizes the format by content
(magic bytes first, extension second) and normalizes to text before the
pipeline runs. The NLP layer then sees the same clean text regardless of
the source format."""
from __future__ import annotations
import io
import re
import zipfile

_TAG = re.compile(r"<[^>]+>")


def _docx_text(data: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        xml = z.read("word/document.xml").decode("utf-8", "replace")
    xml = re.sub(r"<w:tab[^>]*/>", "\t", xml)
    xml = re.sub(r"</w:p>", "\n", xml)
    text = "\n".join(l.strip() for l in _TAG.sub("", xml).split("\n"))
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader
    r = PdfReader(io.BytesIO(data))
    return "\n\n".join((p.extract_text() or "") for p in r.pages).strip()


def _html_text(data: bytes) -> str:
    t = data.decode("utf-8", "replace")
    t = re.sub(r"(?is)<(script|style)[^>]*>.*?</\1>", " ", t)
    t = re.sub(r"(?i)</(p|div|h[1-6]|li|tr)>", "\n", t)
    return re.sub(r"\n{3,}", "\n\n", _TAG.sub("", t)).strip()


def _xlsx_text(data: bytes) -> str:
    """Spreadsheet BRDs/user-story trackers: every sheet becomes a heading,
    every row a line (non-empty cells joined) — the NLP layer then reads
    requirement IDs, stories and modality exactly as in prose."""
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"# Sheet: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None
                     and str(c).strip()]
            if cells:
                parts.append(" | ".join(cells))
    wb.close()
    return re.sub(r"\n{3,}", "\n\n", "\n".join(parts)).strip()


def to_text(filename: str, data: bytes) -> tuple[str, str]:
    """-> (normalized_text, detected_format). Magic bytes beat extensions:
    a .docx renamed .txt still reads as docx."""
    name = (filename or "").lower()
    if data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        raise ValueError("legacy Office 97-2003 format (.doc/.xls) — please "
                         "save as .docx/.xlsx and re-upload")
    if data[:4] == b"PK\x03\x04" and (b"xl/workbook.xml" in data
                                        or name.endswith(".xlsx")):
        try:
            return _xlsx_text(data), "xlsx"
        except Exception:
            pass
    if data[:4] == b"PK\x03\x04" and (b"word/document.xml" in data
                                        or name.endswith(".docx")):
        try:
            return _docx_text(data), "docx"
        except Exception:
            pass
    if data[:5] == b"%PDF-":
        return _pdf_text(data), "pdf"
    if name.endswith((".html", ".htm")) or data[:200].lstrip()[:1] == b"<":
        low = data[:400].lower()
        if b"<html" in low or b"<!doctype" in low or name.endswith((".html", ".htm")):
            return _html_text(data), "html"
    return data.decode("utf-8", "replace"), "text"
