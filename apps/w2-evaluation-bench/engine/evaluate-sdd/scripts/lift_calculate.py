#!/usr/bin/env python3
"""
lift_calculate.py — Section F helper.

Reads two populated score sheets (Output A, Output B) and computes:
- Methodology lift on rounded final scores (A-minus-B labelled; sign reassigned at reveal)
- Per-dim means, gaps to 80% threshold, lift
- Underperformance flags (from Output A perspective; recomputed at reveal)
- Gate pass/fail per lane
- Estimation Handoff Status per lane (OH §3.9)
- Prioritised SA Review List per lane (OH §3.9)

Output is JSON to stdout and optionally to file. Section G consumes via lane_reveal_apply.py.

CRITICAL: this script does NOT reveal lane identity. Output uses A/B labels only.

Usage:
    python lift_calculate.py \\
        --output-a-score-sheet /pilot-runs/<run_id>/output-a-score-sheet.xlsx \\
        --output-b-score-sheet /pilot-runs/<run_id>/output-b-score-sheet.xlsx \\
        [--integration-heavy] \\
        [--confidence-high-risk-a] \\
        [--confidence-high-risk-b] \\
        [--output /pilot-runs/<run_id>/lift-calc.json]
"""

from __future__ import annotations

import argparse
import json
import statistics
import sys
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    print("ERROR: pip install openpyxl --break-system-packages", file=sys.stderr)
    sys.exit(10)

# contracts.py lives in this same scripts/ directory (single source of truth).
sys.path.insert(0, str(Path(__file__).resolve().parent))
import contracts


# (max_base, threshold_base, max_int_heavy, threshold_int_heavy) — DERIVED from
# the single-sourced contracts.DIM_MAX and the 80% gate fraction (v4.1), so the
# maxima and the gate threshold can never drift from the rubric or the scorer.
DIM_WEIGHTS = {
    d: (contracts.dim_max(d, False), contracts.gate_threshold(d, False),
        contracts.dim_max(d, True),  contracts.gate_threshold(d, True))
    for d in range(1, 8)
}

DIM_TO_VARIANCE_ROW = {1: 5, 2: 6, 3: 7, 4: 8, 5: 9, 6: 10, 7: 11}  # v3 Variance_Record data rows (header row 4)
DIM_TO_PERDIM_ROW = {1: 5, 2: 6, 3: 7, 4: 8, 5: 9, 6: 10, 7: 11}

UNDERPERFORMANCE_NEGATIVE_LIFT = -1.0


def get_max_threshold(dim: int, integration_heavy: bool) -> tuple[int, float]:
    base_max, base_thr, inth_max, inth_thr = DIM_WEIGHTS[dim]
    return (inth_max, inth_thr) if integration_heavy else (base_max, base_thr)


def extract_score_sheet(path: Path) -> dict:
    """Read the populated score sheet. Re-derive means/stddevs from the raw five-run
    scores in 6_Variance_Record (the authoritative inputs)."""
    wb = load_workbook(path, data_only=False)

    variance_ws = None
    perdim_ws = None
    deductions_ws = None
    for name in wb.sheetnames:
        nl = name.lower()
        if nl.startswith("variance"):
            variance_ws = wb[name]
        elif nl.startswith("per_dimension"):
            perdim_ws = wb[name]
        elif nl.startswith("deductions") and "audit" not in nl:
            deductions_ws = wb[name]

    if not all([variance_ws, perdim_ws, deductions_ws]):
        raise ValueError(f"{path}: missing required sheets (2/5/6)")

    per_dim_runs = {}
    per_dim_mean = {}
    per_dim_stddev = {}
    per_dim_variance_flag = {}
    for dim, row in DIM_TO_VARIANCE_ROW.items():
        scores = []
        for col in range(3, 8):  # v3 Variance_Record passes in columns C-G
            v = variance_ws.cell(row=row, column=col).value
            if isinstance(v, (int, float)):
                scores.append(float(v))
        per_dim_runs[dim] = scores
        if len(scores) == 5:
            mean = round(sum(scores) / 5, 1)
            stddev = round(statistics.stdev(scores), 2)
            per_dim_mean[dim] = mean
            per_dim_stddev[dim] = stddev
            per_dim_variance_flag[dim] = stddev > 1.0
        else:
            per_dim_mean[dim] = None
            per_dim_stddev[dim] = None
            per_dim_variance_flag[dim] = None

    per_dim_band = {}
    for dim, row in DIM_TO_PERDIM_ROW.items():
        per_dim_band[dim] = perdim_ws.cell(row=row, column=7).value  # v3: band in column G

    means_present = [m for m in per_dim_mean.values() if isinstance(m, (int, float))]
    final_score = int(Decimal(str(sum(means_present))).quantize(Decimal('1'), rounding=ROUND_HALF_UP)) if len(means_present) == 7 else None

    # v3 Deductions sheet is a CATALOGUE: every TRUST/RR row is pre-listed with its
    # standard value in col C. A row counts as an APPLIED deduction only when its
    # triggering-passage cell (col D) has been filled (non-empty, not an unfilled token).
    deductions = []
    for row_idx in range(5, deductions_ws.max_row + 1):
        ded_id = deductions_ws.cell(row=row_idx, column=1).value
        if not isinstance(ded_id, str):
            continue
        ded_id = ded_id.strip()
        if not (ded_id.startswith("TRUST-") or ded_id.startswith("RR-")):
            continue
        passage = deductions_ws.cell(row=row_idx, column=4).value  # col D
        if not isinstance(passage, str) or not passage.strip() or "{{" in passage:
            continue  # catalogue row, not triggered on this run
        deductions.append({
            "id": ded_id,
            "type": deductions_ws.cell(row=row_idx, column=2).value,       # col B
            "deduction": deductions_ws.cell(row=row_idx, column=3).value,  # col C (value)
            "triggering_passage": passage,                                 # col D
            "source": deductions_ws.cell(row=row_idx, column=5).value,     # col E
        })

    return {
        "path": str(path),
        "per_dim_runs": {str(k): v for k, v in per_dim_runs.items()},
        "per_dim_mean": {str(k): v for k, v in per_dim_mean.items()},
        "per_dim_stddev": {str(k): v for k, v in per_dim_stddev.items()},
        "per_dim_variance_flag": {str(k): v for k, v in per_dim_variance_flag.items()},
        "per_dim_band": {str(k): v for k, v in per_dim_band.items()},
        "final_score": final_score,
        "deductions": deductions,
    }


def derive_priority_review(extracted: dict, integration_heavy: bool,
                            confidence_high_risk: bool) -> list[dict]:
    """OH §3.9 mechanical derivation."""
    entries: list[dict] = []

    for dim in range(1, 8):
        _, threshold = get_max_threshold(dim, integration_heavy)
        mean = extracted["per_dim_mean"].get(str(dim))
        if mean is not None and mean < threshold:
            entries.append({
                "priority": "High",
                "section": f"Dim {dim} (sub-threshold)",
                "issue_type": "Gate fail",
                "source_dim": dim,
                "notes": f"Per-dim mean {mean} below 80% threshold {threshold}",
            })

    for d in extracted["deductions"]:
        if isinstance(d.get("id"), str) and d["id"].startswith("TRUST-"):
            entries.append({
                "priority": "High",
                "section": "(see deduction citation)",
                "issue_type": "Trust deduction",
                "source_dim": 3,
                "notes": d.get("triggering_passage") or "",
            })

    for dim in range(1, 8):
        if extracted["per_dim_variance_flag"].get(str(dim)):
            stddev = extracted["per_dim_stddev"].get(str(dim))
            if stddev is None:
                # Dual-judge protocol (v4.7): the flag carries a recorded
                # cross-judge DISSENT, not pass stddev (errata Q6). Same
                # OH §3.9 slot, honest v4.7 wording.
                entries.append({
                    "priority": "Medium",
                    "section": f"Dim {dim} (judge dissent)",
                    "issue_type": "Judge dissent",
                    "source_dim": dim,
                    "notes": ("a recorded cross-judge dissent touches this "
                              "dimension; the conservative consensus stands — "
                              "see the Dissent & Reconciliation annex"),
                })
            else:
                entries.append({
                    "priority": "Medium",
                    "section": f"Dim {dim} (variance)",
                    "issue_type": "Variance",
                    "source_dim": dim,
                    "notes": f"stddev = {stddev} > 1.0",
                })

    for d in extracted["deductions"]:
        if isinstance(d.get("id"), str) and d["id"].startswith("RR-"):
            entries.append({
                "priority": "Medium",
                "section": "(see deduction citation)",
                "issue_type": "Release-awareness deduction",
                "source_dim": 3,
                "notes": d.get("triggering_passage") or "",
            })

    if confidence_high_risk:
        entries.append({
            "priority": "Low",
            "section": "(see Confidence Annotations component)",
            "issue_type": "Confidence flag",
            "source_dim": 4,
            "notes": "Confidence Annotations identified a High-risk section",
        })

    priority_order = {"High": 0, "Medium": 1, "Low": 2}
    entries.sort(key=lambda e: (priority_order.get(e["priority"], 99), e.get("source_dim", 99)))
    return entries


def derive_estimation_handoff_status(extracted: dict, integration_heavy: bool) -> dict:
    """OH §3.9 + deployment-specific resolution of the variance-set edge case."""
    _, dim7_threshold = get_max_threshold(7, integration_heavy)
    dim7_mean = extracted["per_dim_mean"].get("7")
    dim7_variance = extracted["per_dim_variance_flag"].get("7") or False

    if dim7_mean is None:
        return {"status": "Unknown", "basis": "Dim 7 mean unavailable"}

    if dim7_mean < dim7_threshold:
        return {
            "status": "Not Ready",
            "basis": f"Dim 7 mean {dim7_mean} below 80% threshold {dim7_threshold}",
        }

    gate_pass = True
    failing_dims = []
    for dim in range(1, 8):
        _, thr = get_max_threshold(dim, integration_heavy)
        m = extracted["per_dim_mean"].get(str(dim))
        if m is None or m < thr:
            gate_pass = False
            failing_dims.append(dim)

    if gate_pass:
        if dim7_variance:
            # Per OH §3.9 Table 29: Ready requires Dim 7 ≥ threshold AND overall gate
            # cleared AND Dim 7 variance flag NOT set. When Dim 7 ≥ threshold AND gate
            # clears AND variance flag IS set, the status downgrades to Conditional
            # with the variance flag surfaced in basis text for the methodology lead.
            # Dual-judge (v4.7): the flag carries a recorded cross-judge dissent
            # (errata Q6) — same Table-29 downgrade, honest wording.
            dual_judge = extracted["per_dim_stddev"].get("7") is None
            signal = ("a recorded cross-judge dissent touches Dim 7 (conservative "
                      "consensus applied — see the Dissent & Reconciliation annex)"
                      if dual_judge else
                      "Dim 7 variance flag is set (stddev > 1.0)")
            flag_name = "dissent flag" if dual_judge else "variance flag"
            return {
                "status": "Conditional",
                "basis": (
                    f"Dim 7 mean {dim7_mean} ≥ threshold {dim7_threshold} and overall gate cleared, "
                    f"but {signal}. Ready requires the {flag_name} "
                    f"NOT set per OH §3.9 Table 29; the run is downgraded to Conditional and the "
                    f"signal is surfaced for review."
                ),
            }
        return {
            "status": "Ready",
            "basis": (
                f"Dim 7 mean {dim7_mean} ≥ threshold {dim7_threshold}, overall gate cleared, "
                f"Dim 7 variance flag not set"
            ),
        }

    return {
        "status": "Conditional",
        "basis": (
            f"Dim 7 mean {dim7_mean} ≥ threshold {dim7_threshold}, "
            f"but overall gate failed on non-Dim-7 dim(s): {failing_dims}"
        ),
    }


def derive_overall_gate(extracted: dict, integration_heavy: bool) -> dict:
    """Per-dimension gate rolled up to the overall, aligned 1:1 with the four
    bands (rubric s.8): PASS=STRONG (>=80%), MARGINAL_PASS=GOOD (70-79%),
    MARGINAL_FAIL=ADEQUATE (65-69%), FAIL=WEAK (<65%). Build-ready overall needs
    every dimension at STRONG; a GOOD dimension makes the result a conditional
    (accept-with-changes); any ADEQUATE/WEAK dimension fails the gate. The
    variance flag is a soft 'confirm this dimension' annotation and does not move
    the band (single-sourced in contracts.gate_for_pct).
    """
    failing = []   # ADEQUATE / WEAK dims — the gate is not cleared
    marginal = []  # GOOD dims — accept with changes
    for dim in range(1, 8):
        mx, thr = get_max_threshold(dim, integration_heavy)
        m = extracted["per_dim_mean"].get(str(dim))
        vf = extracted.get("per_dim_variance_flag", {}).get(str(dim), False)
        if m is None:
            failing.append({"dim": dim, "mean": m, "threshold": thr, "gate": "FAIL"})
            continue
        g = contracts.gate_for_pct(100.0 * m / mx)
        if g in ("MARGINAL_FAIL", "FAIL"):
            failing.append({"dim": dim, "mean": m, "threshold": thr, "gate": g, "variance_flag": vf})
        elif g == "MARGINAL_PASS":
            marginal.append({"dim": dim, "mean": m, "threshold": thr, "gate": g, "variance_flag": vf})
    return {
        "pass": len(failing) == 0 and len(marginal) == 0,
        "conditional_pass": len(failing) == 0 and len(marginal) > 0,
        "failing_dims": failing,
        "marginal_dims": marginal,
    }


def compute_lift_metrics(extracted_a: dict, extracted_b: dict,
                          integration_heavy: bool) -> dict:
    """Per-dim gaps each lane; per-dim lift A-minus-B; underperformance flags (from A);
       headline lift. Sign reassignment happens in lane_reveal_apply.py."""
    per_dim_gap_a = {}
    per_dim_gap_b = {}
    per_dim_lift = {}
    underperformance_flags_a = {}

    for dim in range(1, 8):
        _, thr = get_max_threshold(dim, integration_heavy)
        m_a = extracted_a["per_dim_mean"].get(str(dim))
        m_b = extracted_b["per_dim_mean"].get(str(dim))

        per_dim_gap_a[str(dim)] = round(m_a - thr, 1) if isinstance(m_a, (int, float)) else None
        per_dim_gap_b[str(dim)] = round(m_b - thr, 1) if isinstance(m_b, (int, float)) else None

        if isinstance(m_a, (int, float)) and isinstance(m_b, (int, float)):
            lift = round(m_a - m_b, 1)
        else:
            lift = None
        per_dim_lift[str(dim)] = lift

        # Underperformance from Output A perspective (recomputed at lane reveal)
        flag = False
        if isinstance(lift, (int, float)) and lift <= UNDERPERFORMANCE_NEGATIVE_LIFT:
            flag = True
        if isinstance(m_a, (int, float)) and m_a < thr:
            flag = True
        underperformance_flags_a[str(dim)] = flag

    headline_lift = None
    if isinstance(extracted_a["final_score"], int) and isinstance(extracted_b["final_score"], int):
        headline_lift = extracted_a["final_score"] - extracted_b["final_score"]

    return {
        "per_dim_gap_output_a": per_dim_gap_a,
        "per_dim_gap_output_b": per_dim_gap_b,
        "per_dim_lift_output_a_minus_b": per_dim_lift,
        "underperformance_flags_output_a_as_candidate_za": underperformance_flags_a,
        "headline_lift_output_a_minus_b": headline_lift,
        "output_a_final_score": extracted_a["final_score"],
        "output_b_final_score": extracted_b["final_score"],
    }


def lift_uncertainty(extracted_a: dict, extracted_b: dict,
                     n_boot: int = 2000, seed: int = 20240101) -> dict | None:
    """Propagate pass-to-pass variability into the headline lift (v4.1).

    The headline lift is a difference of two totals, each a sum of seven means of
    five scoring passes. A bare point estimate ("+3") hides whether that +3 is
    robust to pass noise. This bootstraps the five-pass arrays already stored in
    the score sheet: for each replicate it resamples the five passes per dimension
    (with replacement), re-sums each lane's total, and takes the lift. The result
    is a pass-noise SENSITIVITY BAND (not a frequentist confidence interval) plus
    the fraction of replicates with lift > 0 — an honest statement of how much the
    headline depends on pass variability, holding deductions/floor as scored.
    Because the five passes are produced within one context they are correlated,
    so this band is a LOWER BOUND on true uncertainty, never a tight CI.

    Deterministic (fixed seed) so the interval is reproducible across runs, the
    same reproducibility property the frozen calibration gives the point score.
    Oriented A-minus-B (blinding holds); lane reveal flips the sign with
    flip_uncertainty(). Returns None if any dimension lacks a five-pass array.
    """
    import random
    rng = random.Random(seed)
    runs_a = extracted_a.get("per_dim_runs", {}) or {}
    runs_b = extracted_b.get("per_dim_runs", {}) or {}
    dims = [str(d) for d in range(1, 8)]

    def total_sample(runs):
        tot = 0.0
        for d in dims:
            arr = runs.get(d) or []
            if len(arr) < 2:
                return None
            resample = [arr[rng.randrange(len(arr))] for _ in range(len(arr))]
            tot += sum(resample) / len(resample)
        return tot

    lifts = []
    for _ in range(n_boot):
        ta = total_sample(runs_a)
        tb = total_sample(runs_b)
        if ta is None or tb is None:
            return None
        lifts.append(ta - tb)
    if not lifts:
        return None
    lifts.sort()
    n = len(lifts)

    def pct(p):
        idx = min(n - 1, max(0, int(round(p * (n - 1)))))
        return lifts[idx]

    point = sum(lifts) / n
    p_pos = sum(1 for x in lifts if x > 0) / n

    # ---- Tier 1+2: correlation-aware widening (design effect) ----
    # The naive bootstrap above resamples the five passes AS IF independent, so its
    # band is the LOWER BOUND we always warned about. Estimate ρ̂ across the passes
    # (ICC) and inflate the band's half-width by sqrt(design_effect) around the
    # point estimate, turning the lower-bound band into a correlation-aware one.
    pass_corr = None
    try:
        from pass_correlation import assess_pass_correlation
        # pool both lanes' per-dim arrays — ρ̂ is a property of the pass process,
        # estimated across all dimensions/lanes available.
        pooled = {}
        for lane, runs in (("a", runs_a), ("b", runs_b)):
            for d, arr in (runs or {}).items():
                if isinstance(arr, (list, tuple)) and len(arr) >= 2:
                    pooled[f"{lane}{d}"] = list(arr)
        pass_corr = assess_pass_correlation(pooled)
    except Exception:
        pass_corr = None

    naive_low, naive_high = pct(0.025), pct(0.975)
    if pass_corr and pass_corr.get("reported_se_understatement_factor", 1.0) > 1.0:
        f = pass_corr["reported_se_understatement_factor"]
        adj_low = point + (naive_low - point) * f
        adj_high = point + (naive_high - point) * f
        # P(lift>0) recomputed under a normal approx with the inflated SE, so the
        # probability statement is consistent with the widened band.
        import math as _m
        naive_sd = (naive_high - naive_low) / (2 * 1.96) if naive_high > naive_low else 0.0
        adj_sd = naive_sd * f
        if adj_sd > 0:
            z = point / adj_sd
            p_pos_adj = 0.5 * (1 + _m.erf(z / _m.sqrt(2)))
        else:
            p_pos_adj = p_pos
    else:
        adj_low, adj_high, p_pos_adj = naive_low, naive_high, p_pos

    return {
        "orientation": "output_a_minus_output_b",
        "n_boot": n_boot,
        "seed": seed,
        "lift_mean_bootstrap": round(point, 2),
        # Canonical band = the CORRELATION-AWARE (design-effect-adjusted) band.
        "pass_noise_band_low": round(adj_low, 2),
        "pass_noise_band_high": round(adj_high, 2),
        # The raw, uncorrected bootstrap band, retained for transparency/audit.
        "naive_band_low": round(naive_low, 2),
        "naive_band_high": round(naive_high, 2),
        # Back-compat aliases (older report code reads ci95_*). Now point to the
        # correlation-aware band so consumers get the honest interval by default.
        "lift_ci95_low": round(adj_low, 2),
        "lift_ci95_high": round(adj_high, 2),
        "p_lift_gt_0": round(p_pos_adj, 3),
        "p_lift_gt_0_naive": round(p_pos, 3),
        "pass_correlation": pass_corr,
        "method": ("pass-noise sensitivity band over the five within-context scoring "
                   "passes, deductions/floor held as scored, WIDENED by the design "
                   "effect sqrt(1+(n-1)·ρ̂) estimated via ICC(1) so it reflects the "
                   "passes' correlation rather than assuming independence."),
        "uncertainty_caveat": (
            "This band now accounts for pass-to-pass CORRELATION (via the ICC design "
            "effect), so it is no longer a naive lower bound from assuming the five "
            "passes are independent. It still does NOT capture systematic model bias "
            "(an error shared by all five passes is invisible here), prompt/calibration "
            "sensitivity, or cross-session variance. Re-running the same SDD in a fresh "
            "session may yield a different total. Treat the lift as directional and "
            "prefer band-level distinctions over point scores."),
    }


# ===========================================================================
# v4.7 DUAL-JUDGE PATH (TR-19, Backend Schema §9, errata V-2/C-1): lift is
# computed from the two lanes' SCORING BUNDLES, not by screen-scraping the
# populated XLSX — the sheet is presentation-only under dual-judge. The
# headline is the consensus lift; per-judge lifts (each judge's own A-total
# minus B-total) form the uncertainty band [min, max] of the three. The
# bootstrap/ICC machinery above is BYPASSED: with n=2 judges you report the
# spread, you don't model it. The v4.6 five-pass path is frozen verbatim.
# ===========================================================================

def _dissent_dims_of(bundle: dict) -> set:
    dims = set()
    for d in (bundle.get("dissents") or []):
        cid = d.get("criterion_id") or ""
        sk = d.get("sub_key") or ""
        if cid:
            dims.add(cid[0])
        elif sk.startswith("sub:"):
            dims.add(sk.split(":")[1])
        elif d.get("dim"):
            dims.add(str(d["dim"]))
    return dims


def extract_from_bundle_v47(bundle: dict) -> dict:
    """Build the same 'extracted' shape the derivation helpers consume, from a
    v4.7 scoring bundle. per_dim_variance_flag carries the DISSENT flag
    (stddev None marks the dual-judge protocol for the wording switches)."""
    means = {str(d): bundle.get("per_dim_mean", {}).get(str(d)) for d in range(1, 8)}
    bands = {str(d): bundle.get("per_dim_band", {}).get(str(d)) for d in range(1, 8)}
    dissent_dims = _dissent_dims_of(bundle)
    judge_runs = bundle.get("judge_runs_by_dimension", {}) or {}

    means_present = [m for m in means.values() if isinstance(m, (int, float))]
    final_score = (int(Decimal(str(sum(means_present))).quantize(
        Decimal('1'), rounding=ROUND_HALF_UP)) if len(means_present) == 7 else None)

    judge_totals = {}
    for judge in contracts.JUDGES:
        vals = [(judge_runs.get(judge) or {}).get(str(d)) for d in range(1, 8)]
        judge_totals[judge] = (int(Decimal(str(sum(vals))).quantize(
            Decimal('1'), rounding=ROUND_HALF_UP))
            if all(isinstance(v, (int, float)) for v in vals) else None)

    return {
        "path": "(scoring bundle, dual-judge)",
        "per_dim_runs": {},
        "per_dim_mean": means,
        "per_dim_stddev": {str(d): None for d in range(1, 8)},
        "per_dim_variance_flag": {str(d): (str(d) in dissent_dims) for d in range(1, 8)},
        "per_dim_band": bands,
        "final_score": final_score,
        "judge_totals": judge_totals,
        "deductions": bundle.get("deductions", []) or [],
        "agreement_stats": bundle.get("agreement_stats", {}) or {},
    }


def compute_judge_lifts(extracted_a: dict, extracted_b: dict,
                        headline_consensus) -> tuple[dict, list]:
    """Per-judge lifts (judge's own A-total minus B-total, blinded orientation)
    and the band = [min, max] over {lift_cc, lift_gm, lift_consensus}."""
    judge_lifts = {}
    for judge in contracts.JUDGES:
        ta = (extracted_a.get("judge_totals") or {}).get(judge)
        tb = (extracted_b.get("judge_totals") or {}).get(judge)
        judge_lifts[judge] = (ta - tb) if isinstance(ta, int) and isinstance(tb, int) else None
    vals = [v for v in list(judge_lifts.values()) + [headline_consensus]
            if isinstance(v, (int, float))]
    band = [min(vals), max(vals)] if vals else [None, None]
    return judge_lifts, band


def judge_spread_uncertainty(judge_lifts: dict, band: list,
                             agreement_overall) -> dict:
    """The v4.7 replacement for the bootstrap block (TR-19): report the
    inter-judge spread, don't model it."""
    return {
        "orientation": "output_a_minus_output_b",
        "judge_lifts": dict(judge_lifts),
        "band_low": band[0], "band_high": band[1],
        "agreement_overall": agreement_overall,
        "method": ("inter-judge spread: the band spans the two judges' fully "
                   "independent lifts and the consensus lift. No bootstrap or "
                   "ICC modeling — with two judges the honest statement is the "
                   "spread itself (TR-19)."),
        "uncertainty_caveat": (
            "The band spans the two judges' independent reads; a lift whose "
            "sign holds across both model families is directionally robust. "
            "It does not capture errors shared by both models, prompt/"
            "calibration sensitivity, or cross-session variance. Treat the "
            "lift as directional and prefer band-level distinctions over "
            "point scores."),
    }


def flip_judge_metrics(lift_metrics: dict) -> dict:
    """Reorient the v4.7 judge metrics at lane reveal (B is ZA): negate each
    judge lift; the band negates AND swaps ends (mirror of flip_uncertainty)."""
    lm = dict(lift_metrics)
    jl = lm.get("judge_lifts_output_a_minus_b")
    if isinstance(jl, dict):
        lm["judge_lifts_output_a_minus_b"] = {
            k: (-v if isinstance(v, (int, float)) else v) for k, v in jl.items()}
    band = lm.get("lift_band_output_a_minus_b")
    if isinstance(band, list) and len(band) == 2 and \
            all(isinstance(x, (int, float)) for x in band):
        lm["lift_band_output_a_minus_b"] = [-band[1], -band[0]]
    return lm


def run_agreement_overall(extracted_a: dict, extracted_b: dict):
    """Run-level agreement_overall: non-NA-criteria-weighted mean of the two
    lanes' overalls (errata Q2)."""
    tot, acc = 0, 0.0
    for ex in (extracted_a, extracted_b):
        st = ex.get("agreement_stats") or {}
        o, n = st.get("agreement_overall"), st.get("non_na_criteria")
        if isinstance(o, (int, float)) and isinstance(n, (int, float)) and n:
            acc += o * n
            tot += n
    return round(acc / tot, 3) if tot else None


def flip_uncertainty(u: dict | None) -> dict | None:
    """Reorient an A-minus-B lift-uncertainty block to the revealed sign.
    Negates the mean, swaps and negates the band bounds, and flips P(lift>0).
    Carries both the canonical pass_noise_band_* names and the ci95_* aliases.

    v4.7: also handles the dual-judge judge-spread block (judge_lifts +
    band_low/band_high): each judge lift negates; the band negates and swaps."""
    if not u:
        return u
    if "judge_lifts" in u:  # v4.7 judge-spread block
        flipped = dict(u)
        flipped["orientation"] = "revealed (sign-reassigned at lane reveal)"
        flipped["judge_lifts"] = {k: (-v if isinstance(v, (int, float)) else v)
                                  for k, v in (u.get("judge_lifts") or {}).items()}
        lo, hi = u.get("band_low"), u.get("band_high")
        if isinstance(lo, (int, float)) and isinstance(hi, (int, float)):
            flipped["band_low"], flipped["band_high"] = -hi, -lo
        return flipped
    flipped = dict(u)
    flipped["orientation"] = "revealed (sign-reassigned at lane reveal)"
    flipped["lift_mean_bootstrap"] = round(-u["lift_mean_bootstrap"], 2)
    lo = u.get("pass_noise_band_low", u.get("lift_ci95_low"))
    hi = u.get("pass_noise_band_high", u.get("lift_ci95_high"))
    flipped["pass_noise_band_low"] = round(-hi, 2)
    flipped["pass_noise_band_high"] = round(-lo, 2)
    flipped["lift_ci95_low"] = round(-hi, 2)
    flipped["lift_ci95_high"] = round(-lo, 2)
    flipped["p_lift_gt_0"] = round(1.0 - u["p_lift_gt_0"], 3)
    # New v4.3 fields: flip the naive band + naive p; pass_correlation is
    # sign-agnostic (it describes the pass process) so it carries unchanged.
    if "naive_band_low" in u and "naive_band_high" in u:
        flipped["naive_band_low"] = round(-u["naive_band_high"], 2)
        flipped["naive_band_high"] = round(-u["naive_band_low"], 2)
    if "p_lift_gt_0_naive" in u:
        flipped["p_lift_gt_0_naive"] = round(1.0 - u["p_lift_gt_0_naive"], 3)
    return flipped


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--output-a-score-sheet", type=Path,
                    help="five-pass path input (v4.6, frozen)")
    ap.add_argument("--output-b-score-sheet", type=Path,
                    help="five-pass path input (v4.6, frozen)")
    ap.add_argument("--bundle-a", type=Path,
                    help="dual-judge path input: Output A scoring bundle (v4.7)")
    ap.add_argument("--bundle-b", type=Path,
                    help="dual-judge path input: Output B scoring bundle (v4.7)")
    ap.add_argument("--protocol", choices=["auto", "dual-judge", "five-pass"],
                    default="auto",
                    help="auto: dual-judge when --bundle-a/--bundle-b supplied, "
                         "else five-pass from the score sheets")
    ap.add_argument("--integration-heavy", action="store_true")
    ap.add_argument("--confidence-high-risk-a", action="store_true")
    ap.add_argument("--confidence-high-risk-b", action="store_true")
    ap.add_argument("--output", type=Path)
    args = ap.parse_args()

    protocol = args.protocol
    if protocol == "auto":
        protocol = "dual-judge" if (args.bundle_a and args.bundle_b) else "five-pass"

    if protocol == "dual-judge":
        if not (args.bundle_a and args.bundle_b):
            print("ERROR: dual-judge lift requires --bundle-a and --bundle-b", file=sys.stderr)
            return 12
        bundle_a = json.loads(args.bundle_a.read_text(encoding="utf-8"))
        bundle_b = json.loads(args.bundle_b.read_text(encoding="utf-8"))
        extracted_a = extract_from_bundle_v47(bundle_a)
        extracted_b = extract_from_bundle_v47(bundle_b)
        lift = compute_lift_metrics(extracted_a, extracted_b, args.integration_heavy)
        judge_lifts, band = compute_judge_lifts(
            extracted_a, extracted_b, lift["headline_lift_output_a_minus_b"])
        agreement_overall = run_agreement_overall(extracted_a, extracted_b)
        lift["judge_lifts_output_a_minus_b"] = judge_lifts
        lift["lift_band_output_a_minus_b"] = band
        # Bootstrap/ICC bypassed (TR-19): the uncertainty block IS the spread.
        lift["lift_uncertainty"] = judge_spread_uncertainty(
            judge_lifts, band, agreement_overall)
    else:
        if not (args.output_a_score_sheet and args.output_b_score_sheet):
            print("ERROR: five-pass lift requires --output-a-score-sheet and "
                  "--output-b-score-sheet", file=sys.stderr)
            return 12
        extracted_a = extract_score_sheet(args.output_a_score_sheet)
        extracted_b = extract_score_sheet(args.output_b_score_sheet)
        lift = compute_lift_metrics(extracted_a, extracted_b, args.integration_heavy)
        lift["lift_uncertainty"] = lift_uncertainty(extracted_a, extracted_b)
    gate_a = derive_overall_gate(extracted_a, args.integration_heavy)
    gate_b = derive_overall_gate(extracted_b, args.integration_heavy)
    handoff_a = derive_estimation_handoff_status(extracted_a, args.integration_heavy)
    handoff_b = derive_estimation_handoff_status(extracted_b, args.integration_heavy)
    review_a = derive_priority_review(extracted_a, args.integration_heavy, args.confidence_high_risk_a)
    review_b = derive_priority_review(extracted_b, args.integration_heavy, args.confidence_high_risk_b)

    result = {
        "protocol": protocol,
        "integration_heavy": args.integration_heavy,
        "output_a": {
            "per_dim_mean": extracted_a["per_dim_mean"],
            "per_dim_stddev": extracted_a["per_dim_stddev"],
            "per_dim_variance_flag": extracted_a["per_dim_variance_flag"],
            "per_dim_band": extracted_a["per_dim_band"],
            "final_score": extracted_a["final_score"],
            "deduction_summary": {
                "trust_count": sum(1 for d in extracted_a["deductions"]
                                    if str(d.get("id", "")).startswith("TRUST-")),
                "release_awareness_count": sum(1 for d in extracted_a["deductions"]
                                                if str(d.get("id", "")).startswith("RR-")),
            },
        },
        "output_b": {
            "per_dim_mean": extracted_b["per_dim_mean"],
            "per_dim_stddev": extracted_b["per_dim_stddev"],
            "per_dim_variance_flag": extracted_b["per_dim_variance_flag"],
            "per_dim_band": extracted_b["per_dim_band"],
            "final_score": extracted_b["final_score"],
            "deduction_summary": {
                "trust_count": sum(1 for d in extracted_b["deductions"]
                                    if str(d.get("id", "")).startswith("TRUST-")),
                "release_awareness_count": sum(1 for d in extracted_b["deductions"]
                                                if str(d.get("id", "")).startswith("RR-")),
            },
        },
        "lift_metrics": lift,
        **({"judge_totals_output_a": extracted_a.get("judge_totals"),
            "judge_totals_output_b": extracted_b.get("judge_totals"),
            "agreement_overall": run_agreement_overall(extracted_a, extracted_b)}
           if protocol == "dual-judge" else {}),
        "gate_pass_output_a": gate_a,
        "gate_pass_output_b": gate_b,
        "estimation_handoff_status_output_a": handoff_a,
        "estimation_handoff_status_output_b": handoff_b,
        "priority_review_list_output_a": review_a,
        "priority_review_list_output_b": review_b,
        "note": (
            "All values use Output A / Output B labels. Lane reveal "
            "(remapping to ZA/OTS and sign reassignment) happens in lane_reveal_apply.py."
        ),
    }

    serialised = json.dumps(result, indent=2, default=str)
    print(serialised)
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(serialised, encoding="utf-8")
    return 0


if __name__ == "__main__":
    sys.exit(main())
