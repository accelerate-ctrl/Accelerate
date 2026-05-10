"""XLSX exports for catalogue / lifecycle / clients.

Per spec §17 / ARCHITECTURE Batch 8.

Builds in-memory ``.xlsx`` workbooks via openpyxl (already in
requirements.txt for catalogue ingest). Each export returns ``bytes``
suitable for an HTTP attachment response.

Available exports:

    catalogue.xlsx   pillars / categories / l1 / subcaps with personas
    lifecycle.xlsx   per-subcap lifecycle scores + signals
    clients.xlsx     per-client journey w/ touched subcaps + vendor stack
    benchmarks.xlsx  per-cohort distributions + verdict
"""

from __future__ import annotations

import io
import logging
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .repository import get_repository

logger = logging.getLogger(__name__)

# Brand-aligned header tone for exports.
HEADER_FILL = PatternFill(start_color="103D33", end_color="103D33", fill_type="solid")
HEADER_FONT = Font(color="F1F6EE", bold=True)


def _write_headers(ws, headers: list[str]) -> None:
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=col, value=h)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = Alignment(horizontal="left", vertical="center")
    ws.freeze_panes = "A2"


def _autosize(ws, headers: list[str]) -> None:
    for col_idx in range(1, len(headers) + 1):
        max_len = len(headers[col_idx - 1])
        for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
            for cell in row:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(60, max_len + 2)


def _save(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ─── Builders ───────────────────────────────────────────────────────────────


def export_catalogue() -> bytes:
    repo = get_repository()
    wb = Workbook()
    ws = wb.active
    ws.title = "subcaps"
    headers = [
        "sub_cap_id", "sub_cap_name", "pillar_id", "category_id",
        "l1_capability", "tier", "solution_type", "personas", "description",
    ]
    _write_headers(ws, headers)
    for row, sc in enumerate(sorted(repo.list("subcaps"), key=lambda s: s.get("sub_cap_id", "")), start=2):
        ws.cell(row=row, column=1, value=sc.get("sub_cap_id"))
        ws.cell(row=row, column=2, value=sc.get("sub_cap_name"))
        ws.cell(row=row, column=3, value=sc.get("pillar_id"))
        ws.cell(row=row, column=4, value=sc.get("category_id"))
        ws.cell(row=row, column=5, value=sc.get("l1_capability"))
        ws.cell(row=row, column=6, value=sc.get("tier"))
        ws.cell(row=row, column=7, value=sc.get("solution_type"))
        ws.cell(row=row, column=8, value=", ".join(sc.get("personas") or []))
        ws.cell(row=row, column=9, value=(sc.get("description") or "")[:500])
    _autosize(ws, headers)

    # Categories sheet
    ws_c = wb.create_sheet("categories")
    cat_headers = ["category_id", "name", "pillar_id", "subcap_count"]
    _write_headers(ws_c, cat_headers)
    cats = sorted(repo.list("categories"), key=lambda c: c.get("category_id", ""))
    sub_count_by_cat: dict[str, int] = {}
    for sc in repo.list("subcaps"):
        cat_id = sc.get("category_id")
        if cat_id:
            sub_count_by_cat[cat_id] = sub_count_by_cat.get(cat_id, 0) + 1
    for row, c in enumerate(cats, start=2):
        ws_c.cell(row=row, column=1, value=c.get("category_id"))
        ws_c.cell(row=row, column=2, value=c.get("name"))
        ws_c.cell(row=row, column=3, value=c.get("pillar_id"))
        ws_c.cell(row=row, column=4, value=sub_count_by_cat.get(c.get("category_id"), 0))
    _autosize(ws_c, cat_headers)

    return _save(wb)


def export_lifecycle() -> bytes:
    repo = get_repository()
    wb = Workbook()
    ws = wb.active
    ws.title = "lifecycle"
    headers = [
        "sub_cap_id", "sub_cap_name", "state", "score", "confidence",
        "sow_active", "sow_prospect", "news_last_90d", "benchmark_full",
        "benchmark_indicative", "canonical_stories", "last_signal_at",
    ]
    _write_headers(ws, headers)
    rows = sorted(
        repo.list("lifecycle_scores"),
        key=lambda s: -float(s.get("score") or 0),
    )
    for row, s in enumerate(rows, start=2):
        sig = s.get("signals") or {}
        ws.cell(row=row, column=1, value=s.get("sub_cap_id"))
        ws.cell(row=row, column=2, value=s.get("sub_cap_name"))
        ws.cell(row=row, column=3, value=s.get("state"))
        ws.cell(row=row, column=4, value=s.get("score"))
        ws.cell(row=row, column=5, value=s.get("confidence"))
        ws.cell(row=row, column=6, value=sig.get("sow_active", 0))
        ws.cell(row=row, column=7, value=sig.get("sow_prospect", 0))
        ws.cell(row=row, column=8, value=sig.get("news_last_90d", 0))
        ws.cell(row=row, column=9, value=sig.get("benchmark_full", 0))
        ws.cell(row=row, column=10, value=sig.get("benchmark_indicative", 0))
        ws.cell(row=row, column=11, value=sig.get("canonical_stories", 0))
        ws.cell(row=row, column=12, value=s.get("last_signal_at"))
    _autosize(ws, headers)
    return _save(wb)


def export_clients() -> bytes:
    repo = get_repository()
    wb = Workbook()
    ws = wb.active
    ws.title = "clients"
    headers = [
        "client_name", "active_sows", "prospect_sows", "inactive_sows",
        "subverticals", "cohorts", "asset_size_usd_bn", "touched_subcaps",
        "vendor_count",
    ]
    _write_headers(ws, headers)
    journeys = sorted(repo.list("client_journeys"), key=lambda j: j.get("client_name", ""))
    for row, j in enumerate(journeys, start=2):
        ws.cell(row=row, column=1, value=j.get("client_name"))
        ws.cell(row=row, column=2, value=j.get("sow_count_active", 0))
        ws.cell(row=row, column=3, value=j.get("sow_count_prospect", 0))
        ws.cell(row=row, column=4, value=j.get("sow_count_inactive", 0))
        ws.cell(row=row, column=5, value=", ".join(j.get("subverticals") or []))
        ws.cell(row=row, column=6, value=", ".join(j.get("cohorts") or []))
        ws.cell(row=row, column=7, value=j.get("asset_size_usd_bn"))
        ws.cell(row=row, column=8, value=len(j.get("touched_subcaps") or []))
        ws.cell(row=row, column=9, value=len(j.get("vendor_stack") or []))
    _autosize(ws, headers)

    # Detail sheet — flat list of (client × touched subcap)
    ws_d = wb.create_sheet("client_subcaps")
    detail_headers = ["client_name", "sub_cap_id", "sub_cap_name", "state", "score", "sow_count"]
    _write_headers(ws_d, detail_headers)
    row = 2
    for j in journeys:
        for t in j.get("touched_subcaps") or []:
            ws_d.cell(row=row, column=1, value=j.get("client_name"))
            ws_d.cell(row=row, column=2, value=t.get("sub_cap_id"))
            ws_d.cell(row=row, column=3, value=t.get("sub_cap_name"))
            ws_d.cell(row=row, column=4, value=t.get("state"))
            ws_d.cell(row=row, column=5, value=t.get("score"))
            ws_d.cell(row=row, column=6, value=t.get("sow_count"))
            row += 1
    _autosize(ws_d, detail_headers)

    return _save(wb)


def export_benchmarks() -> bytes:
    repo = get_repository()
    wb = Workbook()
    ws = wb.active
    ws.title = "benchmarks"
    headers = [
        "metric_id", "cohort_id", "period", "n", "p25", "p50", "p75",
        "mean", "stdev", "coef_var", "verdict", "source_kinds",
    ]
    _write_headers(ws, headers)
    rows = sorted(
        repo.list("benchmark_distributions"),
        key=lambda d: (d.get("metric_id", ""), d.get("cohort_id", ""), d.get("period", "")),
    )
    for row, r in enumerate(rows, start=2):
        ws.cell(row=row, column=1, value=r.get("metric_id"))
        ws.cell(row=row, column=2, value=r.get("cohort_id"))
        ws.cell(row=row, column=3, value=r.get("period"))
        ws.cell(row=row, column=4, value=r.get("n"))
        ws.cell(row=row, column=5, value=r.get("p25"))
        ws.cell(row=row, column=6, value=r.get("p50"))
        ws.cell(row=row, column=7, value=r.get("p75"))
        ws.cell(row=row, column=8, value=r.get("mean"))
        ws.cell(row=row, column=9, value=r.get("stdev"))
        ws.cell(row=row, column=10, value=r.get("coef_var"))
        ws.cell(row=row, column=11, value=r.get("verdict"))
        ws.cell(row=row, column=12, value=", ".join(r.get("source_kinds") or []))
    _autosize(ws, headers)
    return _save(wb)
