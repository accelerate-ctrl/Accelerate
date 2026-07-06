#!/usr/bin/env python3
"""report_style.py - shared Zennify-branded docx styling toolkit.

Single source of the brand palette, DM Sans font tokens, and the low-level
python-docx helpers (runs, paragraphs, headings, tables, callouts, header/footer)
used by the report builders. Extracted so report builders share ONE styling
vocabulary and no builder depends on another builder for chrome.
"""
import os
from docx.shared import Pt, RGBColor, Twips, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DARK_TEAL = RGBColor(0x1c, 0x4a, 0x4d)
BRAND_TEAL = RGBColor(0x13, 0x9f, 0x94)
MUTED_BLUE = RGBColor(0x80, 0x94, 0xc0)
SUB_TEAL = RGBColor(0x22, 0xbc, 0xad)
WHITE = RGBColor(0xff, 0xff, 0xff)
GRAY = RGBColor(0x66, 0x66, 0x66)
POS = RGBColor(0x05, 0x96, 0x69)   # positive lift
NEG = RGBColor(0xc2, 0x50, 0x08)   # negative lift
HDR_FILL = "139f94"; ROW_ODD = "f2f4f9"; ROW_EVEN = "e8f7f6"
BAND_FILL = "1c4a4d"; BLUF_FILL = "e8f7f6"
F_REG = "DM Sans"; F_SEMI = "DM Sans SemiBold"; F_MED = "DM Sans Medium"

ASSETS_CANDIDATES = [
    "/mnt/skills/organization/zennify-docx/assets",
    "/mnt/skills/user/zennify-docx/assets",
    "/mnt/skills/private/zennify-docx/assets",
]
ASSETS = next((p for p in ASSETS_CANDIDATES if os.path.isdir(p)), None)
LOGO = os.path.join(ASSETS, "zennify-logo.png") if ASSETS else None
ICON = os.path.join(ASSETS, "zennify-icon.png") if ASSETS else None
HAVE_LOGO = bool(LOGO and os.path.exists(LOGO))
HAVE_ICON = bool(ICON and os.path.exists(ICON))

DIM_MAX_DEFAULT = {1: 15, 2: 15, 3: 20, 4: 15, 5: 10, 6: 10, 7: 15}

_TCPR_ORDER = ["cnfStyle","tcW","gridSpan","hMerge","vMerge","tcBorders","shd","noWrap",
               "tcMar","textDirection","tcFitText","vAlign","hideMark","headers"]
def _tcpr_insert(tcPr, elem):
    tag = elem.tag.split('}')[-1]
    idx = _TCPR_ORDER.index(tag) if tag in _TCPR_ORDER else len(_TCPR_ORDER)
    ins = None
    for child in tcPr:
        ct = child.tag.split('}')[-1]
        ci = _TCPR_ORDER.index(ct) if ct in _TCPR_ORDER else len(_TCPR_ORDER)
        if ci > idx: ins = child; break
    for child in list(tcPr):
        if child.tag == elem.tag: tcPr.remove(child)
    (ins.addprevious(elem) if ins is not None else tcPr.append(elem))

def shade(cell, hexfill):
    e = OxmlElement("w:shd"); e.set(qn("w:val"),"clear"); e.set(qn("w:color"),"auto"); e.set(qn("w:fill"),hexfill)
    _tcpr_insert(cell._tc.get_or_add_tcPr(), e)

def borders(cell, color="ffffff", sz=6, nil=False):
    b = OxmlElement("w:tcBorders")
    for edge in ("top","left","bottom","right"):
        x = OxmlElement(f"w:{edge}")
        if nil: x.set(qn("w:val"),"nil")
        else:
            x.set(qn("w:val"),"single"); x.set(qn("w:sz"),str(sz)); x.set(qn("w:space"),"0"); x.set(qn("w:color"),color)
        b.append(x)
    _tcpr_insert(cell._tc.get_or_add_tcPr(), b)

def margins(cell, t=70, b=70, l=120, r=120):
    m = OxmlElement("w:tcMar")
    for edge,val in (("top",t),("left",l),("bottom",b),("right",r)):
        x = OxmlElement(f"w:{edge}"); x.set(qn("w:w"),str(val)); x.set(qn("w:type"),"dxa"); m.append(x)
    _tcpr_insert(cell._tc.get_or_add_tcPr(), m)

def run(p, text, *, font=F_REG, size=10.5, color=DARK_TEAL, bold=False, italic=False):
    text = (text or "").replace("\u2014", "-")
    r = p.add_run(text); r.font.name=font; r.font.size=Pt(size); r.font.color.rgb=color; r.bold=bold; r.italic=italic
    rPr = r._element.get_or_add_rPr(); rf = rPr.find(qn("w:rFonts"))
    if rf is None: rf = OxmlElement("w:rFonts"); rPr.append(rf)
    for a in ("w:ascii","w:hAnsi","w:cs"): rf.set(qn(a), font)
    return r

def para(doc, *, before=0, after=6, line=None, keep=False, align=None):
    p = doc.add_paragraph(); pf = p.paragraph_format
    pf.space_before=Pt(before); pf.space_after=Pt(after); pf.keep_with_next=keep
    if line: pf.line_spacing=line
    if align is not None: p.alignment=align
    return p

def h1(doc, text):
    p = para(doc, before=18, after=6, keep=True); run(p, text, font=F_SEMI, size=20, color=BRAND_TEAL); return p
def h2(doc, text):
    p = para(doc, before=13, after=4, keep=True); run(p, text, font=F_SEMI, size=13, color=DARK_TEAL); return p
def eyebrow(doc, text):
    p = para(doc, before=2, after=2); run(p, text.upper(), font=F_REG, size=9.5, color=MUTED_BLUE, bold=True); return p

def table(doc, headers, widths=None):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=True
    for i,c in enumerate(t.rows[0].cells):
        shade(c, HDR_FILL); borders(c); margins(c); c.vertical_alignment=WD_ALIGN_VERTICAL.CENTER
        p = c.paragraphs[0]; p.paragraph_format.space_after=Pt(0)
        run(p, headers[i], font=F_SEMI, size=9, color=WHITE, bold=True)
    return t

def row(t, cells, idx, *, bolds=None, colors=None):
    cs = t.add_row().cells; fill = ROW_EVEN if idx%2 else ROW_ODD
    for i,c in enumerate(cs):
        shade(c, fill); borders(c); margins(c); c.vertical_alignment=WD_ALIGN_VERTICAL.CENTER
        p = c.paragraphs[0]; p.paragraph_format.space_after=Pt(0)
        b = (bolds or [False]*len(cs))[i]; col = (colors or [DARK_TEAL]*len(cs))[i]
        run(p, str(cells[i]) if i < len(cells) else "", font=(F_SEMI if b else F_REG), size=9, color=col, bold=b)
    return cs

def callout(doc, lines, fill=BLUF_FILL, title=None):
    t = doc.add_table(rows=1, cols=1); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    c = t.rows[0].cells[0]; shade(c, fill); borders(c, nil=True); margins(c, t=200, b=200, l=240, r=240)
    c.paragraphs[0]._element.getparent().remove(c.paragraphs[0]._element)
    if title:
        tp = c.add_paragraph(); tp.paragraph_format.space_after=Pt(4)
        run(tp, title.upper(), font=F_SEMI, size=10, color=BRAND_TEAL, bold=True)
    for i,(txt, *fmt) in enumerate([(l,) if isinstance(l,str) else l for l in lines]):
        lp = c.add_paragraph(); lp.paragraph_format.space_after=Pt(3 if i < len(lines)-1 else 0); lp.paragraph_format.line_spacing=1.3
        run(lp, txt, size=10.5, color=DARK_TEAL)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)
    return t

def horizontal_rule(doc):
    p = para(doc, after=4); pPr = p._p.get_or_add_pPr(); b = OxmlElement("w:pBdr"); bo = OxmlElement("w:bottom")
    bo.set(qn("w:val"),"single"); bo.set(qn("w:sz"),"6"); bo.set(qn("w:space"),"1"); bo.set(qn("w:color"),"d9e2e1")
    b.append(bo); pPr.append(b)


# ---------- cover + chrome ----------
def cover_banner(doc, *, doc_title, subtitle_bits, grid_pairs, page_break=True):
    """The canonical branded cover used by BOTH report builders (Mode A Diagnostic
    Report and Mode B SDD Review Report), so the two deliverables share one cover.

    doc_title    - e.g. "Diagnostic Report" / "SDD Review Report"
    subtitle_bits- list joined by "  .  " e.g. [engagement, "Comparative (Mode A)"]
    grid_pairs   - list of (LABEL, value) rendered two-per-row in a borderless grid
                   (PREPARED BY / ENGAGEMENT / RUN ID / GENERATED / CALIBRATION /
                   EVALUATOR MODEL or SDD ORIGIN / HEADLINE / CONFIDENTIALITY).
    """
    # Title band (teal), full width.
    band = doc.add_table(rows=1, cols=1); band.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = band.rows[0].cells[0]; shade(c, BAND_FILL); borders(c, nil=True); margins(c, t=300, b=300, l=320, r=320)
    c.paragraphs[0]._element.getparent().remove(c.paragraphs[0]._element)
    tp = c.add_paragraph(); tp.paragraph_format.space_after = Pt(3); tp.paragraph_format.line_spacing = Pt(30)
    run(tp, "Salesforce Solutioning Evaluation  ", font=F_SEMI, size=22, color=WHITE)
    run(tp, doc_title, font=F_REG, size=22, color=SUB_TEAL)
    if subtitle_bits:
        sp = c.add_paragraph(); sp.paragraph_format.space_before = Pt(4); sp.paragraph_format.line_spacing = Pt(16)
        run(sp, "  .  ".join(subtitle_bits), font=F_REG, size=12.5, color=SUB_TEAL)
    doc.add_paragraph().paragraph_format.space_after = Pt(10)
    # Meta grid, two pairs per row, borderless.
    mt = doc.add_table(rows=0, cols=2); mt.alignment = WD_TABLE_ALIGNMENT.CENTER
    for a in range(0, len(grid_pairs), 2):
        cells = mt.add_row().cells
        for cell, (lab, val) in zip(cells, grid_pairs[a:a + 2]):
            borders(cell, nil=True); margins(cell, t=80, b=80, l=110, r=110)
            cell.paragraphs[0]._element.getparent().remove(cell.paragraphs[0]._element)
            lp = cell.add_paragraph(); lp.paragraph_format.space_after = Pt(1)
            run(lp, lab.upper(), font=F_SEMI, size=9.5, color=MUTED_BLUE, bold=True)
            vp = cell.add_paragraph(); run(vp, str(val), size=12.5, color=DARK_TEAL)
    if page_break:
        doc.add_page_break()


def add_logo_header(section):
    section.header.is_linked_to_previous = False
    p = section.header.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if HAVE_LOGO: p.add_run().add_picture(LOGO, width=Emu(1500000), height=Emu(359000))
    else: run(p, "zennify", font=F_SEMI, size=13, color=BRAND_TEAL, bold=True)

def add_footer(section):
    section.footer.is_linked_to_previous = False
    p = section.footer.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if HAVE_ICON:
        p.add_run().add_picture(ICON, width=Emu(150000), height=Emu(150000)); run(p, "   ", size=9)
    fb=OxmlElement("w:fldChar"); fb.set(qn("w:fldCharType"),"begin")
    it=OxmlElement("w:instrText"); it.set(qn("xml:space"),"preserve"); it.text=" PAGE "
    fe=OxmlElement("w:fldChar"); fe.set(qn("w:fldCharType"),"end")
    pr=p.add_run(); pr._element.append(fb); pr._element.append(it); pr._element.append(fe)
    pr.font.name=F_REG; pr.font.size=Pt(9); pr.font.color.rgb=DARK_TEAL; pr.bold=True


def _clip(s, n=60):
    s = (s or "").replace("\u2014","-"); return s if len(s) <= n else s[:n-1].rstrip()+"\u2026"

def _patch_zoom(path):
    import zipfile, re, os as _os
    tmp = path + ".tmp"
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp,"w",zipfile.ZIP_DEFLATED) as zout:
        for it in zin.infolist():
            d = zin.read(it.filename)
            if it.filename == "word/settings.xml":
                s = d.decode("utf-8")
                s = re.sub(r'<w:zoom(?![^>]*w:percent)([^>]*)/>', r'<w:zoom\1 w:percent="100"/>', s)
                s = re.sub(r'<w:zoom(?![^>]*w:percent)([^>]*)>', r'<w:zoom\1 w:percent="100">', s)
                d = s.encode("utf-8")
            zout.writestr(it, d)
    _os.replace(tmp, path)

